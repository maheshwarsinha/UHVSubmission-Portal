import os
import io
import csv
from datetime import datetime, timezone
import pytz
import psutil
import logging
from typing import Optional, List, Dict, Any, Tuple
from fastapi import APIRouter, Depends, HTTPException, Query, status, BackgroundTasks
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.api.deps import get_db, get_current_admin
from app.core.config import settings
from app.models import (
    Administrator, Submission, Student, SubmissionFile,
    BackgroundJob, NotificationRecord, SystemSetting, AuditLog
)

logger = logging.getLogger("admin_api")
from app.schemas.admin import MetricsSummaryResponse, ResetDataRequest
from app.schemas.system import SystemSettingsUpdate
from app.services.upload_service import UploadService
from app.core.security import verify_password

# Tasks for retry
from app.tasks.drive_tasks import execute_drive_transfer
try:
    from app.tasks.drive_tasks import upload_file_to_gdrive_task
    from app.tasks.email_tasks import send_submission_emails_task
    CELERY_AVAILABLE = True
except Exception:
    CELERY_AVAILABLE = False

router = APIRouter()

def log_audit(db: Session, admin_id: str, action: str, resource_type: str = None, resource_id: str = None, details: dict = None):
    audit = AuditLog(
        admin_id=admin_id,
        action=action,
        resource_type=resource_type,
        resource_id=resource_id,
        details=details or {}
    )
    db.add(audit)
    db.commit()

_brevo_quota_cache = {
    "last_checked": 0.0,
    "sent_today": 0,
    "daily_limit": 300,
    "remaining_today": 300
}

def get_brevo_quota_metrics(db: Session, now_ist: datetime) -> Dict[str, Any]:
    """
    Connects to Brevo's live account API (cached for 60 seconds) to display real-time
    remaining email credits. Falls back cleanly to local database counts if offline.
    """
    global _brevo_quota_cache
    current_ts = datetime.now().timestamp()
    if current_ts - _brevo_quota_cache["last_checked"] < 60.0:
        return _brevo_quota_cache

    if getattr(settings, "BREVO_API_KEY", None):
        try:
            import httpx
            headers = {"api-key": settings.BREVO_API_KEY}
            res = httpx.get("https://api.brevo.com/v3/account", headers=headers, timeout=3.5)
            if res.status_code == 200:
                data = res.json()
                plans = data.get("plan", [])
                for p in plans:
                    if p.get("creditsType") == "sendLimit":
                        rem = int(p.get("credits", 300))
                        limit = 300
                        sent = max(0, limit - rem)
                        _brevo_quota_cache = {
                            "last_checked": current_ts,
                            "sent_today": sent,
                            "daily_limit": limit,
                            "remaining_today": rem
                        }
                        return _brevo_quota_cache
        except Exception as e:
            logger.warning(f"Could not fetch live Brevo quota ({e}), using DB count")

    # Fallback: Count from database records
    start_of_today_utc = now_ist.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc)
    db_sent = db.query(NotificationRecord).filter(
        NotificationRecord.recipient_type == "STUDENT",
        NotificationRecord.delivery_status == "SENT",
        NotificationRecord.sent_at >= start_of_today_utc
    ).count()
    limit = 300
    rem = max(0, limit - db_sent)
    _brevo_quota_cache = {
        "last_checked": current_ts,
        "sent_today": db_sent,
        "daily_limit": limit,
        "remaining_today": rem
    }
    return _brevo_quota_cache

def reconcile_submission_drive_files(submission: Submission, db: Session) -> bool:
    """
    Checks Google Drive for any unlinked files belonging to this submission.
    If a file with expected_filename already exists with size > 0 in Admin Drive,
    links it immediately and updates upload_status to DRIVE_COMPLETED.
    """
    from app.services.gdrive_service import gdrive_service
    if not (gdrive_service.service and gdrive_service.enabled):
        return False

    updated = False
    for sub_file in submission.files:
        if sub_file.upload_status != "DRIVE_COMPLETED":
            try:
                # Query Admin Google Drive for expected filename
                q = f"name = '{sub_file.expected_filename}' and trashed = false"
                res = gdrive_service.service.files().list(
                    q=q,
                    spaces="drive",
                    fields="files(id, webViewLink, size)"
                ).execute()
                files = [f for f in res.get("files", []) if int(f.get("size", 0)) > 0]
                if files:
                    drive_id = files[0]["id"]
                    web_link = files[0].get("webViewLink") or f"https://drive.google.com/file/d/{drive_id}/view"
                    sub_file.drive_file_id = drive_id
                    sub_file.admin_drive_url = web_link
                    sub_file.upload_status = "DRIVE_COMPLETED"
                    sub_file.verification_status = True
                    sub_file.error_message = None
                    updated = True
                    logger.info(f"Reconciled and linked {sub_file.expected_filename} to Drive ID: {drive_id}")
            except Exception as e:
                logger.warning(f"Error checking drive reconciliation for {sub_file.expected_filename}: {e}")

    if updated:
        all_done = all(f.upload_status == "DRIVE_COMPLETED" for f in submission.files)
        if all_done:
            submission.status = "COMPLETED"
            submission.drive_completed_at = datetime.now(timezone.utc)
        db.commit()
        db.refresh(submission)
    return updated

@router.get("/metrics", response_model=MetricsSummaryResponse)
def get_metrics(
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Computes real-time metrics for the admin overview dashboard.
    """
    total = db.query(Submission).count()
    completed = db.query(Submission).filter(Submission.status == "COMPLETED").count()
    failed = db.query(Submission).filter(Submission.status == "FAILED").count()
    drive_pending = total - (completed + failed)

    on_time = db.query(Submission).filter(Submission.deadline_classification == "ON_TIME").count()
    late = db.query(Submission).filter(Submission.deadline_classification == "LATE").count()

    # Student category counts
    reg_count = db.query(Submission).join(Student).filter(Student.student_category == "Regular").count()
    backlog_count = db.query(Submission).join(Student).filter(Student.student_category == "Backlog").count()

    # Breakdown by section
    sec_rows = db.query(Student.section, func.count(Submission.id)).join(Submission, Student.id == Submission.student_id).group_by(Student.section).all()
    section_breakdown = {s: cnt for s, cnt in sec_rows}

    # Breakdown by date (in IST)
    ist_tz = pytz.timezone(settings.TIMEZONE)
    subs = db.query(Submission.server_received_at).all()
    daily_breakdown: Dict[str, int] = {}
    for s in subs:
        if s.server_received_at:
            dt_ist = s.server_received_at.astimezone(ist_tz)
            date_key = dt_ist.strftime("%Y-%m-%d")
            daily_breakdown[date_key] = daily_breakdown.get(date_key, 0) + 1

    # Disk usage
    upload_dir = UploadService.get_upload_dir()
    disk = psutil.disk_usage(upload_dir)
    disk_free_gb = round(disk.free / (1024 ** 3), 2)
    disk_total_gb = round(disk.total / (1024 ** 3), 2)
    disk_used_pct = round(disk.percent, 1)

    # Student email quota calculation (Live Brevo Account API sync with DB fallback)
    now_ist = datetime.now(ist_tz)
    quota_metrics = get_brevo_quota_metrics(db, now_ist)
    sent_today = quota_metrics["sent_today"]
    daily_limit = quota_metrics["daily_limit"]
    remaining_today = quota_metrics["remaining_today"]

    queued_count = db.query(NotificationRecord).filter(
        NotificationRecord.delivery_status == "QUEUED_DAILY_LIMIT"
    ).count()

    return MetricsSummaryResponse(
        total_submissions=total,
        completed_submissions=completed,
        drive_pending_submissions=max(0, drive_pending),
        failed_submissions=failed,
        on_time_count=on_time,
        late_count=late,
        regular_count=reg_count,
        backlog_count=backlog_count,
        section_breakdown=section_breakdown,
        daily_breakdown=daily_breakdown,
        disk_free_gb=disk_free_gb,
        disk_total_gb=disk_total_gb,
        disk_used_percent=disk_used_pct,
        email_sent_today=sent_today,
        email_daily_limit=daily_limit,
        email_remaining_today=remaining_today,
        email_queued_count=queued_count
    )

@router.get("/submissions")
def get_submissions_list(
    search: Optional[str] = None,
    section: Optional[str] = None,
    branch: Optional[str] = None,
    status: Optional[str] = None,
    category: Optional[str] = None,
    deadline_status: Optional[str] = None,
    sort_by: Optional[str] = "date",
    sort_order: Optional[str] = "desc",
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Search, filter, and sort submissions with pagination.
    """
    query = db.query(Submission).join(Student)

    if search:
        search_pattern = f"%{search.strip()}%"
        query = query.filter(
            (Student.name.ilike(search_pattern)) |
            (Student.roll_number.ilike(search_pattern)) |
            (Student.section.ilike(search_pattern)) |
            (Student.branch.ilike(search_pattern)) |
            (Submission.topic.ilike(search_pattern)) |
            (Submission.tracking_id.ilike(search_pattern))
        )

    if section:
        query = query.filter(Student.section == section)

    if branch:
        query = query.filter(Student.branch == branch)

    if status:
        query = query.filter(Submission.status == status)

    if category:
        query = query.filter(Student.student_category == category)

    if deadline_status:
        query = query.filter(Submission.deadline_classification == deadline_status)

    total_count = query.count()

    # Dynamic sorting
    order_desc = (sort_order.lower() == "desc") if sort_order else True
    if sort_by == "roll":
        query = query.order_by(Student.roll_number.desc() if order_desc else Student.roll_number.asc())
    elif sort_by == "name":
        query = query.order_by(Student.name.desc() if order_desc else Student.name.asc())
    elif sort_by == "section":
        query = query.order_by(Student.section.desc() if order_desc else Student.section.asc())
    elif sort_by == "timing":
        query = query.order_by(Submission.deadline_classification.desc() if order_desc else Submission.deadline_classification.asc())
    else:  # default date
        query = query.order_by(Submission.server_received_at.desc() if order_desc else Submission.server_received_at.asc())

    items = query.offset(skip).limit(limit).all()

    ist_tz = pytz.timezone(settings.TIMEZONE)
    results = []
    for item in items:
        server_received_ist = item.server_received_at.astimezone(ist_tz).strftime("%d %b %Y, %I:%M %p")
        all_drive_done = len(item.files) > 0 and all(f.upload_status == "DRIVE_COMPLETED" for f in item.files)
        has_failed = any(f.upload_status == "FAILED" for f in item.files)
        drive_status = "COMPLETED" if all_drive_done else ("FAILED" if has_failed else "QUEUED")

        results.append({
            "id": item.id,
            "tracking_id": item.tracking_id,
            "student_name": item.student.name,
            "roll_number": item.student.roll_number,
            "section": item.student.section,
            "branch": item.student.branch,
            "category": item.student.student_category,
            "email": item.student.email,
            "ip_address": item.ip_address,
            "topic": item.topic,
            "status": item.status,
            "drive_status": drive_status,
            "deadline_classification": item.deadline_classification,
            "server_received_at_ist": server_received_ist,
            "source_folder_url": item.source_folder_url,
            "source_folder_id": item.source_folder_id,
            "folder_name": item.folder_name,
            "folder_verification_status": item.folder_verification_status,
            "files_count": len(item.files),
            "drive_completed": bool(item.drive_completed_at) or all_drive_done,
            "files": [
                {
                    "id": f.id,
                    "file_type": f.file_type,
                    "expected_filename": f.expected_filename,
                    "upload_status": f.upload_status,
                    "source_drive_url": f.source_drive_url,
                    "source_file_id": f.source_file_id,
                    "link_accessibility_status": f.link_accessibility_status,
                    "drive_file_id": f.drive_file_id,
                    "admin_drive_url": f.admin_drive_url or (f"https://drive.google.com/file/d/{f.drive_file_id}/view" if f.drive_file_id else None)
                } for f in item.files
            ]
        })

    return {
        "total": total_count,
        "items": results,
        "skip": skip,
        "limit": limit
    }

@router.get("/submissions/{submission_id}")
def get_submission_detail(
    submission_id: str,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    submission = db.query(Submission).filter(Submission.id == submission_id).first()
    if not submission:
        raise HTTPException(status_code=404, detail="Submission not found")

    # Auto-reconcile with Google Drive if any files are pending but already uploaded
    reconcile_submission_drive_files(submission, db)

    ist_tz = pytz.timezone(settings.TIMEZONE)
    server_time_ist = submission.server_received_at.astimezone(ist_tz).strftime("%d %B %Y, %I:%M:%S %p IST")
    drive_completed_ist = submission.drive_completed_at.astimezone(ist_tz).strftime("%d %B %Y, %I:%M:%S %p IST") if submission.drive_completed_at else None

    files_data = []
    for f in submission.files:
        files_data.append({
            "id": f.id,
            "file_type": f.file_type,
            "original_filename": f.original_filename,
            "expected_filename": f.expected_filename,
            "file_size": f.file_size,
            "file_size_mb": round(f.file_size / (1024 * 1024), 2) if f.file_size else 0.0,
            "mime_type": f.mime_type,
            "sha256": f.sha256_checksum,
            "source_drive_url": f.source_drive_url,
            "source_file_id": f.source_file_id,
            "link_accessibility_status": f.link_accessibility_status,
            "upload_status": f.upload_status,
            "verification_status": f.verification_status,
            "drive_file_id": f.drive_file_id,
            "admin_drive_url": f.admin_drive_url or (f"https://drive.google.com/file/d/{f.drive_file_id}/view" if f.drive_file_id else None),
            "error_message": f.error_message
        })

    jobs_data = []
    for j in submission.jobs:
        jobs_data.append({
            "id": j.id,
            "job_type": j.job_type,
            "status": j.status,
            "retry_count": j.retry_count,
            "max_retries": j.max_retries,
            "last_error": j.last_error,
            "created_at": j.created_at
        })

    notifications_data = []
    for n in submission.notifications:
        notifications_data.append({
            "id": n.id,
            "recipient_type": n.recipient_type,
            "recipient_address": n.recipient_address,
            "notification_type": n.notification_type,
            "delivery_status": n.delivery_status,
            "sent_at": n.sent_at
        })

    return {
        "id": submission.id,
        "tracking_id": submission.tracking_id,
        "student": {
            "name": submission.student.name,
            "roll_number": submission.student.roll_number,
            "section": submission.student.section,
            "branch": submission.student.branch,
            "academic_year": submission.student.academic_year,
            "category": submission.student.student_category,
            "email": submission.student.email
        },
        "topic": submission.topic,
        "status": submission.status,
        "deadline_classification": submission.deadline_classification,
        "ip_address": submission.ip_address,
        "server_received_at_ist": server_time_ist,
        "drive_completed_at_ist": drive_completed_ist,
        "source_folder_url": submission.source_folder_url,
        "source_folder_id": submission.source_folder_id,
        "folder_name": submission.folder_name,
        "folder_verification_status": submission.folder_verification_status,
        "files": files_data,
        "jobs": jobs_data,
        "notifications": notifications_data
    }

@router.post("/jobs/{job_id}/retry")
def retry_failed_job(
    job_id: str,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    job = db.query(BackgroundJob).filter(BackgroundJob.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Background job not found")

    job.status = "QUEUED"
    job.last_error = None
    db.commit()

    if CELERY_AVAILABLE:
        if job.job_type == "gdrive_upload" and job.file_id:
            upload_file_to_gdrive_task.delay(job.submission_id, job.file_id)
        elif job.job_type in ["email_student", "email_teacher"]:
            send_submission_emails_task.delay(job.submission_id)

    log_audit(
        db,
        admin_id=current_admin.id,
        action="JOB_RETRY",
        resource_type="BackgroundJob",
        resource_id=job.id,
        details={"job_type": job.job_type, "submission_id": job.submission_id}
    )

    return {"success": True, "message": f"Job {job_id} queued for retry."}

@router.get("/export-csv")
def export_submissions_csv(
    branch: Optional[str] = None,
    section: Optional[str] = None,
    category: Optional[str] = None,
    deadline_status: Optional[str] = None,
    search: Optional[str] = None,
    sort_by: Optional[str] = "date",
    sort_order: Optional[str] = "desc",
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Exports submission metadata to CSV format with all applied filters.
    """
    query = db.query(Submission).join(Student)
    if search:
        search_pattern = f"%{search.strip()}%"
        query = query.filter(
            (Student.name.ilike(search_pattern)) |
            (Student.roll_number.ilike(search_pattern)) |
            (Student.section.ilike(search_pattern)) |
            (Student.branch.ilike(search_pattern)) |
            (Submission.topic.ilike(search_pattern)) |
            (Submission.tracking_id.ilike(search_pattern))
        )
    if branch:
        query = query.filter(Student.branch == branch)
    if section:
        query = query.filter(Student.section == section)
    if category:
        query = query.filter(Student.student_category == category)
    if deadline_status:
        query = query.filter(Submission.deadline_classification == deadline_status)

    order_desc = (sort_order.lower() == "desc") if sort_order else True
    if sort_by == "roll":
        query = query.order_by(Student.roll_number.desc() if order_desc else Student.roll_number.asc())
    elif sort_by == "name":
        query = query.order_by(Student.name.desc() if order_desc else Student.name.asc())
    elif sort_by == "section":
        query = query.order_by(Student.section.desc() if order_desc else Student.section.asc())
    elif sort_by == "timing":
        query = query.order_by(Submission.deadline_classification.desc() if order_desc else Submission.deadline_classification.asc())
    else:
        query = query.order_by(Submission.server_received_at.desc() if order_desc else Submission.server_received_at.asc())

    submissions = query.all()

    ist_tz = pytz.timezone(settings.TIMEZONE)
    output = io.StringIO()
    writer = csv.writer(output)

    # Write CSV Header
    writer.writerow([
        "Tracking ID", "Roll Number", "Student Name", "Branch", "Section",
        "Category", "Academic Year", "Email", "IP Address", "Topic", "Deadline Status",
        "Server Received (IST)", "Google Drive Status", "Files Submitted", "Student Folder Link",
        "Video Drive Link", "PPT Drive Link", "PDF Drive Link", "Admin Backup Drive Link", "Link Accessibility"
    ])

    for sub in submissions:
        ist_time = sub.server_received_at.astimezone(ist_tz).strftime("%Y-%m-%d %H:%M:%S IST")
        file_names = "; ".join([f.expected_filename for f in sub.files])
        video_f = next((f for f in sub.files if f.file_type == "video"), None)
        ppt_f = next((f for f in sub.files if f.file_type == "ppt"), None)
        pdf_f = next((f for f in sub.files if f.file_type == "pdf"), None)

        video_link = (video_f.source_drive_url if video_f and video_f.source_drive_url else "") or ""
        ppt_link = (ppt_f.source_drive_url if ppt_f and ppt_f.source_drive_url else "") or ""
        pdf_link = (pdf_f.source_drive_url if pdf_f and pdf_f.source_drive_url else "") or ""
        admin_link = (video_f.admin_drive_url if video_f and video_f.admin_drive_url else (ppt_f.admin_drive_url if ppt_f else "")) or ""
        access_notes = "; ".join([f"{f.file_type}: {f.link_accessibility_status}" for f in sub.files if f.link_accessibility_status]) or "Normal"

        writer.writerow([
            sub.tracking_id,
            sub.student.roll_number,
            sub.student.name,
            sub.student.branch,
            sub.student.section,
            sub.student.student_category,
            sub.student.academic_year,
            sub.student.email,
            sub.ip_address or "Unknown",
            sub.topic,
            sub.deadline_classification,
            ist_time,
            sub.status,
            file_names,
            sub.source_folder_url or "",
            video_link,
            ppt_link,
            pdf_link,
            admin_link,
            access_notes
        ])

    log_audit(
        db,
        admin_id=current_admin.id,
        action="EXPORT_CSV",
        details={
            "section": section,
            "branch": branch,
            "category": category,
            "deadline_status": deadline_status,
            "search": search,
            "records_count": len(submissions)
        }
    )

    output.seek(0)
    filename = f"uhv_submissions_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.get("/export-excel")
def export_submissions_excel(
    branch: Optional[str] = None,
    section: Optional[str] = None,
    category: Optional[str] = None,
    deadline_status: Optional[str] = None,
    search: Optional[str] = None,
    sort_by: Optional[str] = "date",
    sort_order: Optional[str] = "desc",
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Exports submission metadata with all active filters to Microsoft Excel (.xlsx) format.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    query = db.query(Submission).join(Student)
    if search:
        search_pattern = f"%{search.strip()}%"
        query = query.filter(
            (Student.name.ilike(search_pattern)) |
            (Student.roll_number.ilike(search_pattern)) |
            (Student.section.ilike(search_pattern)) |
            (Student.branch.ilike(search_pattern)) |
            (Submission.topic.ilike(search_pattern)) |
            (Submission.tracking_id.ilike(search_pattern))
        )
    if branch:
        query = query.filter(Student.branch == branch)
    if section:
        query = query.filter(Student.section == section)
    if category:
        query = query.filter(Student.student_category == category)
    if deadline_status:
        query = query.filter(Submission.deadline_classification == deadline_status)

    order_desc = (sort_order.lower() == "desc") if sort_order else True
    if sort_by == "roll":
        query = query.order_by(Student.roll_number.desc() if order_desc else Student.roll_number.asc())
    elif sort_by == "name":
        query = query.order_by(Student.name.desc() if order_desc else Student.name.asc())
    elif sort_by == "section":
        query = query.order_by(Student.section.desc() if order_desc else Student.section.asc())
    elif sort_by == "timing":
        query = query.order_by(Submission.deadline_classification.desc() if order_desc else Submission.deadline_classification.asc())
    else:
        query = query.order_by(Submission.server_received_at.desc() if order_desc else Submission.server_received_at.asc())

    submissions = query.all()
    ist_tz = pytz.timezone(settings.TIMEZONE)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Submissions"

    headers = [
        "Tracking ID", "Roll Number", "Student Name", "Branch", "Section",
        "Category", "Academic Year", "Email", "IP Address", "Topic", "Deadline Status",
        "Server Received (IST)", "Google Drive Status", "Files Submitted", "Student Folder Link",
        "Video Drive Link", "PPT Drive Link", "PDF Drive Link", "Admin Backup Drive Link", "Link Accessibility"
    ]
    ws.append(headers)

    # Style Header
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Add Data Rows
    for row_idx, sub in enumerate(submissions, start=2):
        ist_time = sub.server_received_at.astimezone(ist_tz).strftime("%d-%m-%Y %I:%M:%S %p")
        file_names = "; ".join([f.expected_filename for f in sub.files])
        all_drive_done = len(sub.files) > 0 and all(f.upload_status == "DRIVE_COMPLETED" for f in sub.files)
        drive_status = "Uploaded" if all_drive_done else "Queued/Pending"

        video_f = next((f for f in sub.files if f.file_type == "video"), None)
        ppt_f = next((f for f in sub.files if f.file_type == "ppt"), None)
        pdf_f = next((f for f in sub.files if f.file_type == "pdf"), None)

        video_link = (video_f.source_drive_url if video_f and video_f.source_drive_url else "") or ""
        ppt_link = (ppt_f.source_drive_url if ppt_f and ppt_f.source_drive_url else "") or ""
        pdf_link = (pdf_f.source_drive_url if pdf_f and pdf_f.source_drive_url else "") or ""
        admin_link = (video_f.admin_drive_url if video_f and video_f.admin_drive_url else (ppt_f.admin_drive_url if ppt_f else "")) or ""
        access_notes = "; ".join([f"{f.file_type}: {f.link_accessibility_status}" for f in sub.files if f.link_accessibility_status]) or "Normal"

        row_data = [
            sub.tracking_id,
            sub.student.roll_number,
            sub.student.name,
            sub.student.branch,
            sub.student.section,
            sub.student.student_category,
            sub.student.academic_year,
            sub.student.email,
            sub.ip_address or "Unknown",
            sub.topic,
            sub.deadline_classification,
            ist_time,
            drive_status,
            file_names,
            sub.source_folder_url or "",
            video_link,
            ppt_link,
            pdf_link,
            admin_link,
            access_notes
        ]
        ws.append(row_data)

        # Style row cells
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if col_idx in [1, 2, 5, 6, 7, 10, 11, 12]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    log_audit(
        db,
        admin_id=current_admin.id,
        action="EXPORT_EXCEL",
        details={"records_count": len(submissions), "branch": branch, "section": section, "category": category}
    )

    filename = f"uhv_submissions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.post("/submissions/{submission_id}/upload-drive")
def trigger_submission_drive_upload(
    submission_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Triggers Google Drive transfer asynchronously with background status updates.
    Performs instant reconciliation for any files already archived in Google Drive.
    """
    submission = db.query(Submission).filter(Submission.id == submission_id).first()
    if not submission:
        raise HTTPException(status_code=404, detail="Submission not found")

    # 1. Instant reconciliation: check if any files are already present in Google Drive
    reconcile_submission_drive_files(submission, db)

    # 2. Identify remaining uncompleted files
    from app.services.drive_queue import drive_queue
    pending_files = [f for f in submission.files if f.upload_status != "DRIVE_COMPLETED"]
    queued_count = 0

    for sub_file in pending_files:
        if drive_queue.enqueue(submission.id, sub_file.id):
            queued_count += 1

    all_done = all(f.upload_status == "DRIVE_COMPLETED" for f in submission.files)
    if all_done:
        submission.status = "COMPLETED"
        submission.drive_completed_at = datetime.now(timezone.utc)
        db.commit()
        msg = "All files successfully archived in Google Drive!"
    else:
        msg = f"Transfer running in background ({queued_count} files queued in serial queue)."

    log_audit(
        db,
        admin_id=current_admin.id,
        action="FORCE_DRIVE_UPLOAD",
        resource_type="Submission",
        resource_id=submission_id,
        details={"queued": queued_count, "all_completed": all_done}
    )

    return {
        "success": True,
        "message": msg,
        "queued_files": queued_count,
        "all_completed": all_done,
        "submission_status": submission.status
    }

@router.post("/submissions/{submission_id}/transfer-file/{file_id}")
def transfer_single_submission_file(
    submission_id: str,
    file_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Admin endpoint to trigger transfer of a single submission file into the Admin Google Drive.
    Enqueues in the serial DriveTransferQueue to prevent concurrent server overloads.
    """
    submission = db.query(Submission).filter(Submission.id == submission_id).first()
    if not submission:
        raise HTTPException(status_code=404, detail="Submission not found")

    reconcile_submission_drive_files(submission, db)
    sub_file = db.query(SubmissionFile).filter(SubmissionFile.id == file_id).first()
    if not sub_file:
        raise HTTPException(status_code=404, detail="File not found")

    if sub_file.upload_status == "DRIVE_COMPLETED":
        return {
            "status": "completed",
            "drive_id": sub_file.drive_file_id,
            "web_link": sub_file.admin_drive_url,
            "message": "File verified and linked in Google Drive."
        }

    from app.services.drive_queue import drive_queue
    drive_queue.enqueue(submission_id, file_id)
    return {
        "status": "processing",
        "message": "File transfer added to serial queue."
    }

@router.post("/test-email")
def send_test_email(
    target_email: Optional[str] = None,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Sends a test email to verify transactional email delivery via Brevo/Resend/SMTP.
    """
    from app.services.email_service import email_service
    teacher_setting = db.query(SystemSetting).filter(SystemSetting.key == "teacher_email").first()
    dyn_teacher = teacher_setting.value if teacher_setting and teacher_setting.value else None
    recipient = target_email or dyn_teacher or settings.TEACHER_EMAIL or settings.EMAIL_FROM
    if not recipient:
        raise HTTPException(status_code=400, detail="No recipient email address configured.")

    subject = "[UHV Portal] Live Email Delivery Test"
    html_content = f"""
    <div style="font-family: Arial, sans-serif; max-width: 500px; padding: 20px; border: 1px solid #e2e8f0; border-radius: 8px;">
        <h3 style="color: #0284c7; margin-top: 0;">UHV Portal Email Test</h3>
        <p>This is a live test email dispatched from the UHV Assignment Portal administration console.</p>
        <p><b>Timestamp:</b> {email_service.get_ist_time_str()}</p>
        <p><b>Admin:</b> {current_admin.email}</p>
        <p style="color: #16a34a; font-weight: bold;">Transactional email delivery is operating correctly.</p>
    </div>
    """
    text_content = f"UHV Portal Email Test. Sent at {email_service.get_ist_time_str()} by {current_admin.email}."

    success, msg = email_service.send_smtp_email(recipient, subject, html_content, text_content, prefer_provider="resend")

    log_audit(
        db,
        admin_id=current_admin.id,
        action="TEST_EMAIL",
        resource_type="Email",
        resource_id=recipient,
        details={"success": success, "message": msg}
    )

    return {
        "success": success,
        "recipient": recipient,
        "message": msg
    }

@router.post("/submissions/{submission_id}/resend-email")
def resend_submission_emails(
    submission_id: str,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Manually triggers email dispatch (student confirmation & teacher notification) for a submission.
    """
    from app.tasks.email_tasks import execute_email_dispatch
    res = execute_email_dispatch(submission_id)
    return res

@router.post("/dispatch-queued-emails")
def trigger_dispatch_queued_emails(
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Manually triggers dispatch of emails queued due to daily quota limits.
    """
    from app.tasks.email_tasks import dispatch_queued_notifications
    res = dispatch_queued_notifications()
    return res

@router.delete("/reset-all-data")
@router.post("/reset-all-data")
def reset_all_portal_data(
    payload: ResetDataRequest,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Clears all submissions, files, jobs, and notifications from the portal and Google Drive to start completely from zero.
    Requires verified administrator password to ensure zero accidental data loss in production.
    """
    if not payload.password or not verify_password(payload.password, current_admin.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect administrator password! Portal reset has been rejected to protect production submissions."
        )

    import shutil
    from app.services.gdrive_service import gdrive_service

    # 1. Clean Google Drive submissions folder
    drive_deleted = gdrive_service.clean_all_submissions()

    # 2. Clean database records
    db.query(NotificationRecord).delete()
    db.query(BackgroundJob).delete()
    db.query(SubmissionFile).delete()
    db.query(Submission).delete()
    db.query(Student).delete()
    db.commit()

    # 3. Clean local storage
    upload_dir = UploadService.get_upload_dir()
    for item in os.listdir(upload_dir):
        item_path = os.path.join(upload_dir, item)
        try:
            if os.path.isdir(item_path):
                shutil.rmtree(item_path, ignore_errors=True)
            else:
                os.remove(item_path)
        except Exception:
            pass

    # Recreate empty base dirs
    os.makedirs(os.path.join(upload_dir, "chunks"), exist_ok=True)
    os.makedirs(os.path.join(upload_dir, "completed"), exist_ok=True)

    log_audit(
        db,
        admin_id=current_admin.id,
        action="RESET_ALL_DATA",
        details={"drive_items_deleted": drive_deleted}
    )

    return {
        "success": True,
        "message": "All portal submissions and uploaded files have been completely cleared from database, local disk, and Google Drive.",
        "drive_items_deleted": drive_deleted
    }

@router.get("/settings")
def get_system_settings(
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    settings_records = db.query(SystemSetting).all()
    result = {}
    for r in settings_records:
        result[r.key] = r.value

    # Defaults if not customized
    result.setdefault("submission_deadline", settings.SUBMISSION_DEADLINE)
    result.setdefault("submissions_open", settings.SUBMISSIONS_OPEN)
    result.setdefault("allowed_sections", settings.ALLOWED_SECTIONS)
    result.setdefault("max_ppt_size_mb", settings.MAX_PPT_SIZE_MB)
    result.setdefault("max_pdf_size_mb", settings.MAX_PDF_SIZE_MB)
    result.setdefault("max_video_size_mb", settings.MAX_VIDEO_SIZE_MB)
    result.setdefault("pdf_optional_allowed", settings.PDF_OPTIONAL_ALLOWED)
    result.setdefault("pdf_required", settings.PDF_REQUIRED)
    result.setdefault("teacher_email", settings.TEACHER_EMAIL)
    result.setdefault("gdrive_root_folder_name", settings.GDRIVE_ROOT_FOLDER_NAME)

    return result

@router.put("/settings")
def update_system_settings(
    payload: SystemSettingsUpdate,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    updates = payload.model_dump(exclude_unset=True)
    for key, val in updates.items():
        record = db.query(SystemSetting).filter(SystemSetting.key == key).first()
        if not record:
            record = SystemSetting(key=key, value=val)
            db.add(record)
        else:
            record.value = val

    db.commit()

    log_audit(
        db,
        admin_id=current_admin.id,
        action="CONFIG_UPDATE",
        details=updates
    )

    return {"success": True, "message": "System configuration updated successfully."}

from pydantic import BaseModel

class GDriveTokenUpdate(BaseModel):
    refresh_token: str

@router.get("/gdrive/status")
def get_gdrive_status(
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    try:
        from app.services.gdrive_service import gdrive_service
        from app.services.drive_queue import drive_queue

        active_token = gdrive_service.get_active_refresh_token()
        is_valid, msg = gdrive_service.reconnect() if active_token else (False, "No refresh token configured")
        
        pending_count = db.query(SubmissionFile).filter(SubmissionFile.upload_status != "DRIVE_COMPLETED").count()
        masked_token = (active_token[:8] + "..." + active_token[-6:]) if len(active_token) > 14 else ("Configured" if active_token else "Not Set")

        return {
            "connected": is_valid,
            "message": msg,
            "has_token": bool(active_token),
            "masked_token": masked_token,
            "pending_files_count": pending_count,
            "queue_status": drive_queue.get_status()
        }
    except Exception as e:
        logger.error(f"Error in get_gdrive_status: {e}")
        return {
            "connected": False,
            "message": str(e),
            "has_token": False,
            "masked_token": "Error",
            "pending_files_count": 0,
            "queue_status": {"queue_size": 0, "current_task": None, "is_busy": False}
        }

def exchange_google_code_for_token(code: str) -> Tuple[Optional[str], Optional[str]]:
    import urllib.request
    import urllib.parse
    import json
    unquoted = urllib.parse.unquote(code.strip())
    payload = urllib.parse.urlencode({
        "code": unquoted,
        "client_id": settings.GDRIVE_CLIENT_ID,
        "client_secret": settings.GDRIVE_CLIENT_SECRET,
        "redirect_uri": "https://developers.google.com/oauthplayground",
        "grant_type": "authorization_code"
    }).encode("utf-8")
    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=payload, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            rt = data.get("refresh_token")
            if rt:
                return rt, None
            return None, "Google did not return a refresh_token (make sure prompt=consent is used)"
    except urllib.error.HTTPError as he:
        err_msg = he.read().decode("utf-8", errors="ignore")
        return None, f"Google OAuth error: {err_msg}"
    except Exception as ex:
        return None, str(ex)

@router.post("/gdrive/update-token")
def update_gdrive_token(
    payload: GDriveTokenUpdate,
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    raw_token = payload.refresh_token.strip()
    if not raw_token:
        raise HTTPException(status_code=400, detail="Token or code cannot be empty")

    token = raw_token
    # If the user pasted an authorization code (starts with 4/ or 4%2F)
    if raw_token.startswith("4/") or raw_token.startswith("4%2F"):
        exchanged_token, err = exchange_google_code_for_token(raw_token)
        if not exchanged_token:
            raise HTTPException(status_code=400, detail=f"Failed to exchange authorization code: {err}")
        token = exchanged_token

    from app.services.gdrive_service import gdrive_service
    from app.services.drive_queue import drive_queue

    try:
        is_valid, msg = gdrive_service.reconnect(token)
        if not is_valid:
            raise HTTPException(status_code=400, detail=f"Google Drive verification failed: {msg}")

        # Save to SystemSetting in DB
        setting = db.query(SystemSetting).filter(SystemSetting.key == "gdrive_refresh_token").first()
        if not setting:
            setting = SystemSetting(key="gdrive_refresh_token", value=token, description="Admin Google Drive OAuth Refresh Token")
            db.add(setting)
        else:
            setting.value = token
        db.commit()

        # Re-initialize service
        gdrive_service._init_service(token)

        # Automatically enqueue all pending/failed files
        pending_files = db.query(SubmissionFile).filter(SubmissionFile.upload_status != "DRIVE_COMPLETED").all()
        queued_count = 0
        for pf in pending_files:
            if drive_queue.enqueue(pf.submission_id, pf.id):
                queued_count += 1

        log_audit(
            db,
            admin_id=current_admin.id,
            action="GDRIVE_TOKEN_UPDATED",
            details={"queued_files": queued_count}
        )

        return {
            "success": True,
            "message": f"Google Drive connected successfully! {queued_count} pending files queued for background transfer.",
            "queued_files": queued_count
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error updating gdrive token: {e}")
        raise HTTPException(status_code=500, detail=f"Internal error updating token: {e}")

@router.post("/gdrive/sync-all-pending")
def sync_all_pending_drive_files(
    db: Session = Depends(get_db),
    current_admin: Administrator = Depends(get_current_admin)
):
    """
    Enqueues all pending/failed files across all submissions into the serial DriveTransferQueue.
    """
    from app.services.drive_queue import drive_queue

    pending_files = db.query(SubmissionFile).filter(SubmissionFile.upload_status != "DRIVE_COMPLETED").all()
    queued_count = 0
    for pf in pending_files:
        if drive_queue.enqueue(pf.submission_id, pf.id):
            queued_count += 1

    return {
        "success": True,
        "message": f"{queued_count} files added to serial Google Drive transfer queue.",
        "queued_files": queued_count
    }
